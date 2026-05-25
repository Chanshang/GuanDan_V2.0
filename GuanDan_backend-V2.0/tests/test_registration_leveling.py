import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from services.registration_leveling import (
    build_person_scores,
    fill_registration_levels,
    split_member_names,
)


class TestRegistrationLeveling(unittest.TestCase):
    def test_split_member_names_accepts_common_separators(self):
        self.assertEqual(split_member_names("张三-李四"), ["张三", "李四"])
        self.assertEqual(split_member_names(" 张三 、 李四 "), ["张三", "李四"])
        self.assertEqual(split_member_names("张三，李四"), ["张三", "李四"])

    def test_build_person_scores_uses_inverse_final_rank(self):
        ranking_rows = [
            {"rank": 1, "member_name": "甲-乙"},
            {"rank": 2, "member_name": "丙-丁"},
            {"rank": 3, "member_name": "戊-己"},
        ]

        scores = build_person_scores(ranking_rows)

        self.assertEqual(scores["甲"], 3)
        self.assertEqual(scores["乙"], 3)
        self.assertEqual(scores["丙"], 2)
        self.assertEqual(scores["丁"], 2)
        self.assertEqual(scores["戊"], 1)
        self.assertNotIn("新选手", scores)

    def test_fill_registration_levels_writes_abc_by_team_total_score(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmpdir = Path(tmpdir)
            registration_path = tmpdir / "registration.xlsx"
            ranking_path = tmpdir / "ranking.xlsx"
            output_path = tmpdir / "registration_filled.xlsx"

            self._make_registration_workbook(registration_path)
            self._make_ranking_workbook(ranking_path)

            result = fill_registration_levels(
                registration_path=registration_path,
                ranking_path=ranking_path,
                output_path=output_path,
            )

            self.assertEqual(result["team_count"], 3)
            self.assertEqual(result["level_counts"], {"A": 1, "B": 1, "C": 1})

            workbook = load_workbook(output_path)
            sheet = workbook.active
            self.assertEqual(sheet["F3"].value, "B")
            self.assertEqual(sheet["I3"].value, "A")
            self.assertEqual(sheet["L3"].value, "C")

    def _make_registration_workbook(self, path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([None] * 17)
        sheet.append(
            [
                "序号",
                "办公室",
                "队长",
                "队名",
                "小组A",
                "A小组等级",
                "队员",
                "小组B",
                "B小组等级",
                "队员",
                "小组C",
                "C小组等级",
                "队员",
                "小组D",
                "D小组等级",
                "队员",
                "替补队员（2名）",
            ]
        )
        sheet.append(
            [
                1,
                "测试办公室",
                "队长",
                "测试队",
                "测试-A",
                None,
                "甲-新选手",
                "测试-B",
                None,
                "丙-丁",
                "测试-C",
                None,
                "新选手1-新选手2",
            ]
        )
        workbook.save(path)

    def _make_ranking_workbook(self, path):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "office", "team_name", "member_name", "big_score", "small_score", "level"])
        sheet.append([1, "1", "T1", "甲-乙", 6, 20, "A"])
        sheet.append([2, "1", "T2", "丙-丁", 4, 20, "B"])
        sheet.append([3, "1", "T3", "戊-己", 2, 20, "C"])
        workbook.save(path)


if __name__ == "__main__":
    unittest.main()
