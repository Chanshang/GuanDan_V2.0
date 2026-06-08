from __future__ import annotations

import math
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


REGISTRATION_TEAM_COLUMNS = (
    {"team_name": 5, "level": 6, "members": 7},
    {"team_name": 8, "level": 9, "members": 10},
    {"team_name": 11, "level": 12, "members": 13},
    {"team_name": 14, "level": 15, "members": 16},
)

MEMBER_SEPARATORS_PATTERN = re.compile(r"[-－—、，,／/]+")
SPACE_PATTERN = re.compile(r"\s+")


def normalize_member_name(value):
    """统一姓名格式，避免空格差异影响跨表匹配。"""
    if value is None:
        return ""
    return SPACE_PATTERN.sub("", str(value)).strip()


def split_member_names(value):
    """拆分一组双人小队的队员姓名，兼容常见中英文分隔符。"""
    if value is None:
        return []
    text = str(value).strip()
    if not text or text == "空":
        return []
    normalized_text = SPACE_PATTERN.sub("", text)
    return [name for name in MEMBER_SEPARATORS_PATTERN.split(normalized_text) if name]


def build_person_scores(ranking_rows):
    """
    将上一届最终排名转换为个人分。

    计算方式：上一届排名越靠前分越高。若上一届共有 N 条有效排名，
    第 1 名个人分为 N，第 2 名为 N-1，依次递减到 1。
    同一上一届小队的两名队员获得相同个人分；新报名表中没出现于上一届
    排名表的姓名在算总分时按 0 分处理，相当于 C 档基准。
    """
    valid_rows = []
    for index, row in enumerate(ranking_rows, start=1):
        names = split_member_names(row.get("member_name"))
        if not names:
            continue
        rank = _parse_rank(row.get("rank"), fallback=index)
        valid_rows.append({"rank": rank, "names": names})

    ordered_rows = sorted(valid_rows, key=lambda item: item["rank"])
    row_count = len(ordered_rows)
    person_scores = {}
    for order_index, row in enumerate(ordered_rows):
        score = row_count - order_index
        for name in row["names"]:
            normalized_name = normalize_member_name(name)
            if normalized_name:
                person_scores[normalized_name] = max(score, person_scores.get(normalized_name, 0))
    return person_scores


def fill_registration_levels(registration_path, ranking_path, output_path=None):
    registration_path = Path(registration_path)
    ranking_path = Path(ranking_path)
    output_path = Path(output_path) if output_path else _default_output_path(registration_path)

    ranking_rows = read_ranking_rows(ranking_path)
    person_scores = build_person_scores(ranking_rows)

    workbook = load_workbook(registration_path)
    sheet = workbook.active
    teams = collect_registration_teams(sheet, person_scores)
    ranked_teams = rank_registration_teams(teams)

    level_counts = Counter()
    for team in ranked_teams:
        sheet.cell(row=team["row"], column=team["level_col"]).value = team["level"]
        level_counts[team["level"]] += 1

    workbook.save(output_path)
    return {
        "output_path": str(output_path),
        "team_count": len(ranked_teams),
        "level_counts": {level: level_counts.get(level, 0) for level in ("A", "B", "C")},
        "missing_member_count": _count_missing_members(ranked_teams),
        "missing_members": sorted(
            {
                member
                for team in ranked_teams
                for member in team["member_scores"]
                if team["member_scores"][member] == 0
            }
        ),
    }


def read_ranking_rows(ranking_path):
    workbook = load_workbook(ranking_path, data_only=True)
    sheet = workbook.active
    headers = _read_headers(sheet)
    rank_col = headers.get("id", 1)
    member_col = headers.get("member_name", 4)

    rows = []
    for row_index in range(2, sheet.max_row + 1):
        rank_value = sheet.cell(row=row_index, column=rank_col).value
        member_name = sheet.cell(row=row_index, column=member_col).value
        if member_name is None:
            continue
        rows.append({"rank": rank_value, "member_name": member_name})
    return rows


def collect_registration_teams(sheet, person_scores):
    teams = []
    sequence = 0
    for row_index in range(3, sheet.max_row + 1):
        for columns in REGISTRATION_TEAM_COLUMNS:
            member_text = sheet.cell(row=row_index, column=columns["members"]).value
            member_names = split_member_names(member_text)
            if not member_names:
                continue

            team_name = sheet.cell(row=row_index, column=columns["team_name"]).value
            member_scores = {
                normalize_member_name(member): person_scores.get(normalize_member_name(member), 0)
                for member in member_names
            }
            teams.append(
                {
                    "sequence": sequence,
                    "row": row_index,
                    "team_name": team_name,
                    "level_col": columns["level"],
                    "members": member_names,
                    "member_scores": member_scores,
                    "total_score": sum(member_scores.values()),
                }
            )
            sequence += 1
    return teams


def rank_registration_teams(teams):
    ranked_teams = sorted(teams, key=lambda item: (-item["total_score"], item["sequence"]))
    team_count = len(ranked_teams)
    a_cutoff = math.ceil(team_count / 3)
    b_cutoff = math.ceil(team_count * 2 / 3)

    for rank, team in enumerate(ranked_teams, start=1):
        if rank <= a_cutoff:
            team["level"] = "A"
        elif rank <= b_cutoff:
            team["level"] = "B"
        else:
            team["level"] = "C"
        team["rank"] = rank
    return ranked_teams


def _read_headers(sheet):
    headers = {}
    for column_index in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=column_index).value
        if value is None:
            continue
        headers[str(value).strip().lower()] = column_index
    return headers


def _parse_rank(value, fallback):
    try:
        return int(value)
    except (TypeError, ValueError):
        return fallback


def _default_output_path(path):
    return path.with_name(f"{path.stem}-已填等级{path.suffix}")


def _count_missing_members(teams):
    return sum(
        1
        for team in teams
        for member_score in team["member_scores"].values()
        if member_score == 0
    )
